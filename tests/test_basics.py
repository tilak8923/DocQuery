"""Fast tests that need no model downloads.

Run:  .venv/bin/python -m unittest discover tests -v
"""

import io
import unittest
from pathlib import Path

from dlqa.chunker import chunk_blocks, chunk_text
from dlqa.loader import Block, load_bytes, load_document
from dlqa.pipeline import is_follow_up

SAMPLES = Path(__file__).resolve().parent.parent / "data" / "samples"


class ChunkTextTests(unittest.TestCase):
    def test_overlap_and_coverage(self):
        words = [f"w{i}" for i in range(500)]
        chunks = chunk_text(" ".join(words), chunk_words=100, overlap_words=20)
        self.assertEqual(chunks[0].text.split()[0], "w0")
        self.assertEqual(chunks[-1].text.split()[-1], "w499")
        self.assertEqual(chunks[1].start_word, 80)
        self.assertTrue(all(len(c.text.split()) <= 100 for c in chunks))

    def test_bad_overlap_rejected(self):
        with self.assertRaises(ValueError):
            chunk_text("a b c", chunk_words=10, overlap_words=10)


class ChunkBlocksTests(unittest.TestCase):
    def test_sections_and_pages(self):
        blocks = [
            Block("Warranty", 3, "heading"),
            Block("Two years from purchase.", 3),
            Block("Cleaning", 4, "heading"),
            Block("Wipe with a damp cloth.", 4),
        ]
        # Small sections are packed together; the first heading is the prefix,
        # later headings stay inline.
        chunks = chunk_blocks(blocks, "m.pdf", chunk_words=50, overlap_words=10)
        self.assertEqual(len(chunks), 1)
        self.assertEqual(chunks[0].section, "Warranty")
        self.assertEqual(chunks[0].page, 3)
        self.assertTrue(chunks[0].text.startswith("Warranty\n"))
        self.assertIn("\nCleaning\n", chunks[0].text)
        self.assertEqual(chunks[0].doc, "m.pdf")
        # Nearest-heading lookup for citations
        pos = chunks[0].text.index("damp cloth")
        self.assertEqual(chunks[0].section_at(pos), "Cleaning")
        self.assertEqual(chunks[0].section_at(chunks[0].text.index("Two years")), "Warranty")
        # With a tight budget each section gets its own chunk.
        chunks = chunk_blocks(blocks, "m.pdf", chunk_words=6, overlap_words=1)
        self.assertEqual([c.section for c in chunks], ["Warranty", "Cleaning"])
        self.assertEqual(chunks[1].page, 4)

    def test_long_block_is_windowed_with_overlap(self):
        long = Block(" ".join(f"w{i}" for i in range(300)), 1)
        chunks = chunk_blocks([Block("Intro", 1, "heading"), long], chunk_words=100, overlap_words=20)
        self.assertGreaterEqual(len(chunks), 3)
        self.assertTrue(all(c.section == "Intro" for c in chunks))
        self.assertIn("w80", chunks[1].text)  # overlap carried over
        self.assertIn("w299", chunks[-1].text)

    def test_table_rows_not_split(self):
        rows = [Block(f"Name: Item {i}; Price: {i}", 1, "table") for i in range(60)]
        chunks = chunk_blocks(rows, chunk_words=40, overlap_words=5)
        for c in chunks:
            for line in c.text.split("\n"):
                self.assertTrue(line.startswith("Name: Item"))


class LoaderTests(unittest.TestCase):
    def test_txt_structure(self):
        doc = load_document(SAMPLES / "syllabus.txt")
        kinds = {b.kind for b in doc.blocks}
        self.assertIn("heading", kinds)
        self.assertIn("Instructor: Dr. Priya Raman", doc.text)
        self.assertGreater(doc.words, 400)

    def test_pdf_pages(self):
        doc = load_document(SAMPLES / "transformers.pdf")
        self.assertEqual(doc.pages, 1)
        self.assertIn("Attention Is All You Need", doc.text)
        self.assertTrue(all(b.page == 1 for b in doc.blocks))

    def test_docx_headings_and_tables(self):
        import docx

        d = docx.Document()
        d.add_heading("Warranty", level=1)
        d.add_paragraph("Two years from the date of purchase.")
        t = d.add_table(rows=3, cols=2)
        for r, (k, v) in enumerate([("Model", "AB-2000"), ("Power", "2200 W"), ("Weight", "1.2 kg")]):
            t.cell(r, 0).text, t.cell(r, 1).text = k, v
        buf = io.BytesIO()
        d.save(buf)
        doc = load_bytes(buf.getvalue(), "m.docx")
        self.assertEqual(doc.blocks[0].kind, "heading")
        # First table row is treated as the header, so later rows become "Model: Power; AB-2000: 2200 W"
        self.assertTrue(any(b.kind == "table" and "2200 W" in b.text for b in doc.blocks))

    def test_pptx_slides(self):
        from pptx import Presentation
        from pptx.util import Inches

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Results"
        slide.placeholders[1].text = "Accuracy improved to 88 percent."
        slide2 = prs.slides.add_slide(prs.slide_layouts[5])
        slide2.shapes.title.text = "Next steps"
        box = slide2.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
        box.text_frame.text = "Fine-tune the reader."
        buf = io.BytesIO()
        prs.save(buf)
        doc = load_bytes(buf.getvalue(), "deck.pptx")
        self.assertEqual(doc.pages, 2)
        self.assertTrue(any(b.kind == "heading" and b.text == "Results" for b in doc.blocks))
        self.assertTrue(any(b.page == 2 and "Fine-tune" in b.text for b in doc.blocks))

    def test_xlsx_rows_become_sentences(self):
        import pandas as pd

        buf = io.BytesIO()
        pd.DataFrame({"Product": ["Kettle", "Toaster"], "Price": [49, 29]}).to_excel(buf, index=False, sheet_name="Catalog")
        doc = load_bytes(buf.getvalue(), "cat.xlsx")
        self.assertIn("Sheet: Catalog", doc.text)
        self.assertIn("Product: Kettle; Price: 49", doc.text)

    def test_csv(self):
        doc = load_bytes(b"name,age\nAsha,31\nRavi,28\n", "people.csv")
        self.assertIn("name: Asha; age: 31", doc.text)

    def test_html(self):
        html = b"<html><body><h1>Guide</h1><p>Boil water first.</p><ul><li>Step one</li></ul><table><tr><td>A</td><td>B</td></tr></table><script>x()</script></body></html>"
        doc = load_bytes(html, "g.html")
        kinds = [b.kind for b in doc.blocks]
        self.assertEqual(kinds, ["heading", "paragraph", "list", "table"])
        self.assertNotIn("x()", doc.text)

    def test_unsupported_type(self):
        with self.assertRaises(ValueError):
            load_bytes(b"x", "file.exe")


class FollowUpTests(unittest.TestCase):
    def test_detection(self):
        self.assertTrue(is_follow_up("what does it do"))
        self.assertTrue(is_follow_up("how long is that?"))
        self.assertFalse(is_follow_up("What is the capacity of the kettle?"))
        self.assertFalse(is_follow_up("Explain how the self-attention mechanism works in the transformer"))


if __name__ == "__main__":
    unittest.main()
