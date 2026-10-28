"""Document loader: any supported file -> a Document made of structured Blocks.

Supported: PDF, DOCX, PPTX, XLSX/XLS, CSV, HTML, TXT, MD, and images
(PNG/JPG/TIFF/BMP/WEBP) via OCR. Scanned PDFs are OCR'd page by page when
Tesseract is installed; otherwise the Document is flagged `needs_ocr`.

Why blocks instead of one big string: headings, paragraphs and table rows
are kept apart so the chunker can keep each chunk inside one section, carry
the page number for citations, and never cut a table row in half.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED_TYPES = [
    "pdf", "docx", "pptx", "xlsx", "xls", "csv", "html", "htm", "txt", "md",
    "png", "jpg", "jpeg", "tiff", "tif", "bmp", "webp",
]

MAX_TABLE_ROWS = 2000


@dataclass
class Block:
    text: str
    page: int | None = None  # 1-based page / slide / sheet number
    kind: str = "paragraph"  # paragraph | heading | table | list


@dataclass
class Document:
    name: str
    blocks: list[Block] = field(default_factory=list)
    pages: int = 0
    needs_ocr: bool = False  # True when text could not be extracted and OCR was unavailable
    notes: list[str] = field(default_factory=list)  # warnings to show the user (limits hit)

    @property
    def text(self) -> str:
        return "\n\n".join(b.text for b in self.blocks)

    @property
    def words(self) -> int:
        return sum(len(b.text.split()) for b in self.blocks)


# ---------------------------------------------------------------- entry points
def load_document(path: str | Path, progress=None) -> Document:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(path)
    return load_bytes(path.read_bytes(), path.name, progress)


def load_bytes(data: bytes, filename: str, progress=None) -> Document:
    """`progress(done, total)` is called per page for PDFs (used by the UI progress bar)."""
    suffix = Path(filename).suffix.lower().lstrip(".")
    if suffix not in SUPPORTED_TYPES:
        raise ValueError(f"Unsupported file type: .{suffix}")
    loader = {
        "pdf": _load_pdf,
        "docx": _load_docx,
        "pptx": _load_pptx,
        "xlsx": _load_excel,
        "xls": _load_excel,
        "csv": _load_csv,
        "html": _load_html,
        "htm": _load_html,
        "txt": _load_text,
        "md": _load_text,
    }.get(suffix, _load_image)
    doc = _load_pdf(data, filename, progress) if suffix == "pdf" else loader(data, filename)
    doc.blocks = [b for b in doc.blocks if b.text.strip()]
    return doc


# ---------------------------------------------------------------- helpers
def _clean(text: str) -> str:
    text = text.replace("\xa0", " ").replace(" ", " ")
    return " ".join(text.split())


def _row_to_text(header: list[str], row: list[str]) -> str:
    cells = [_clean(str(c)) if c is not None else "" for c in row]
    if header and any(header):
        pairs = [f"{h}: {v}" for h, v in zip(header, cells) if v and h]
        return "; ".join(pairs)
    return " | ".join(c for c in cells if c)


def _table_blocks(rows: list[list], page: int | None) -> list[Block]:
    rows = [r for r in rows if r and any(c not in (None, "") for c in r)]
    if not rows:
        return []
    header = [_clean(str(c)) if c is not None else "" for c in rows[0]]
    out = []
    for r in rows[1:MAX_TABLE_ROWS]:
        t = _row_to_text(header, r)
        if t:
            out.append(Block(t, page, "table"))
    if len(rows) == 1:  # header-only table: keep it as a line
        out.append(Block(" | ".join(h for h in header if h), page, "table"))
    return out


def _looks_like_heading(text: str) -> bool:
    words = text.split()
    if not words or len(words) > 12:
        return False
    if text.endswith((".", ":", ";", ",")):
        return False
    letters = [c for c in text if c.isalpha()]
    if letters and sum(c.isupper() for c in letters) / len(letters) > 0.8:
        return True
    return text.istitle() and len(words) <= 6


def _ocr_image(img) -> str | None:
    """Return OCR text, or None when Tesseract is not available."""
    try:
        import pytesseract

        return pytesseract.image_to_string(img)
    except Exception:  # noqa: BLE001  (missing binary, missing package)
        return None


# ---------------------------------------------------------------- PDF
def _page_has_rulings(page) -> bool:
    """find_tables() (default "lines" strategy) only finds tables drawn with vector
    lines/rectangles, and costs ~150 ms per page. Pages with no drawings (novels,
    plain notes) cannot contain such a table, so we skip the expensive call there."""
    try:
        get = getattr(page, "get_cdrawings", None) or page.get_drawings
        return len(get()) > 0
    except Exception:  # noqa: BLE001
        return True  # unknown -> be safe and look for tables


def _load_pdf(data: bytes, filename: str, progress=None) -> Document:
    import pymupdf

    from .config import settings

    doc = Document(name=filename)
    needs_ocr = False
    ocr_pages = 0
    with pymupdf.open(stream=data, filetype="pdf") as pdf:
        doc.pages = len(pdf)
        detect_tables = len(pdf) <= settings.table_detect_max_pages
        if not detect_tables:
            doc.notes.append(
                f"{len(pdf)} pages: table detection skipped (limit {settings.table_detect_max_pages} pages)."
            )
        for pno, page in enumerate(pdf, start=1):
            if progress is not None:
                progress(pno - 1, doc.pages)
            table_rects = []
            # Tables first, so their text is not duplicated as paragraphs.
            try:
                tables = page.find_tables().tables if detect_tables and _page_has_rulings(page) else []
                for tab in tables:
                    table_rects.append(pymupdf.Rect(tab.bbox))
                    doc.blocks.extend(_table_blocks(tab.extract(), pno))
            except Exception:  # noqa: BLE001  (older PyMuPDF without find_tables)
                pass

            info = page.get_text("dict")
            sizes = [
                span["size"]
                for b in info["blocks"] if b.get("type") == 0
                for line in b["lines"] for span in line["spans"] if span["text"].strip()
            ]
            body_size = sorted(sizes)[len(sizes) // 2] if sizes else 0

            page_has_text = False
            for b in info["blocks"]:
                if b.get("type") != 0:
                    continue
                rect = pymupdf.Rect(b["bbox"])
                if any(rect.intersects(tr) for tr in table_rects):
                    continue
                lines = [
                    "".join(span["text"] for span in line["spans"]) for line in b["lines"]
                ]
                text = _clean(" ".join(lines))
                if not text:
                    continue
                page_has_text = True
                max_size = max(
                    (span["size"] for line in b["lines"] for span in line["spans"]), default=0
                )
                is_heading = (body_size and max_size >= body_size * 1.15 and len(text.split()) <= 12) or _looks_like_heading(text)
                doc.blocks.append(Block(text, pno, "heading" if is_heading else "paragraph"))

            if not page_has_text and not table_rects:
                # Probably a scanned page: try OCR (capped, it costs seconds per page).
                if ocr_pages >= settings.max_ocr_pages:
                    if ocr_pages == settings.max_ocr_pages:
                        doc.notes.append(
                            f"Only the first {settings.max_ocr_pages} scanned pages were OCR'd "
                            "(DLQA_MAX_OCR_PAGES); the rest were skipped."
                        )
                        ocr_pages += 1
                    continue
                ocr_pages += 1
                pix = page.get_pixmap(dpi=150)
                try:
                    from PIL import Image

                    img = Image.open(io.BytesIO(pix.tobytes("png")))
                    text = _ocr_image(img)
                except Exception:  # noqa: BLE001
                    text = None
                if text is None:
                    needs_ocr = True
                else:
                    doc.blocks.extend(_text_to_blocks(text, pno))
    doc.needs_ocr = needs_ocr and not doc.blocks
    return doc


# ---------------------------------------------------------------- DOCX
def _load_docx(data: bytes, filename: str) -> Document:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(io.BytesIO(data))
    doc = Document(name=filename)
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, d)
            text = _clean(p.text)
            if not text:
                continue
            style = (p.style.name or "").lower() if p.style is not None else ""
            if style.startswith(("heading", "title")):
                kind = "heading"
            elif "list" in style:
                kind = "list"
            else:
                kind = "paragraph"
            doc.blocks.append(Block(text, None, kind))
        elif tag == "tbl":
            t = Table(child, d)
            rows = [[c.text for c in r.cells] for r in t.rows]
            doc.blocks.extend(_table_blocks(rows, None))
    return doc


# ---------------------------------------------------------------- PPTX
def _load_pptx(data: bytes, filename: str) -> Document:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    doc = Document(name=filename)
    for sno, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        if title_shape is not None and title_shape.has_text_frame and title_shape.text_frame.text.strip():
            doc.blocks.append(Block(_clean(title_shape.text_frame.text), sno, "heading"))
        for shape in slide.shapes:
            if shape is title_shape:
                continue
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[c.text for c in r.cells] for r in shape.table.rows]
                doc.blocks.extend(_table_blocks(rows, sno))
            elif shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = _clean("".join(r.text for r in para.runs))
                    if text:
                        doc.blocks.append(Block(text, sno, "list" if para.level else "paragraph"))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = _clean(slide.notes_slide.notes_text_frame.text)
            if notes:
                doc.blocks.append(Block(f"Speaker notes: {notes}", sno, "paragraph"))
    doc.pages = len(prs.slides)
    return doc


# ---------------------------------------------------------------- spreadsheets
def _load_excel(data: bytes, filename: str) -> Document:
    import pandas as pd

    sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, header=None, dtype=str)
    doc = Document(name=filename)
    for sno, (sheet, df) in enumerate(sheets.items(), start=1):
        doc.blocks.append(Block(f"Sheet: {sheet}", sno, "heading"))
        rows = df.fillna("").values.tolist()
        doc.blocks.extend(_table_blocks(rows, sno))
    doc.pages = len(sheets)
    return doc


def _load_csv(data: bytes, filename: str) -> Document:
    import pandas as pd

    df = pd.read_csv(io.BytesIO(data), header=None, dtype=str, on_bad_lines="skip")
    doc = Document(name=filename, pages=1)
    rows = df.fillna("").values.tolist()
    doc.blocks.extend(_table_blocks(rows, 1))
    return doc


# ---------------------------------------------------------------- HTML
def _load_html(data: bytes, filename: str) -> Document:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(data, "html.parser")
    for tag in soup(["script", "style", "noscript", "nav", "footer"]):
        tag.decompose()
    doc = Document(name=filename, pages=1)
    for el in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "tr", "pre", "blockquote"]):
        if el.name == "tr":
            cells = [_clean(c.get_text(" ")) for c in el.find_all(["td", "th"])]
            text = " | ".join(c for c in cells if c)
            kind = "table"
        else:
            if el.find_parent(["li", "tr", "p"]) is not None and el.name in ("p", "li"):
                continue  # nested: parent already captured it
            text = _clean(el.get_text(" "))
            kind = "heading" if el.name[0] == "h" else ("list" if el.name == "li" else "paragraph")
        if text:
            doc.blocks.append(Block(text, 1, kind))
    if not doc.blocks:
        doc.blocks.extend(_text_to_blocks(soup.get_text("\n"), 1))
    return doc


# ---------------------------------------------------------------- text / markdown
def _load_text(data: bytes, filename: str) -> Document:
    text = data.decode("utf-8", errors="ignore")
    doc = Document(name=filename, pages=1)
    doc.blocks.extend(_text_to_blocks(text, 1))
    return doc


_LIST_RE = re.compile(r"^\s*([-*•]|\d+[.)])\s+")


def _is_title_line(line: str, next_line: str) -> bool:
    """A short capitalised line with no end punctuation, followed by a line that
    starts a new sentence, is a heading glued to its paragraph."""
    words = line.split()
    if not (1 <= len(words) <= 6) or not line[0].isupper():
        return False
    if ":" in line or line[-1] in ".,;:!?)":
        return False
    return bool(next_line) and (next_line[0].isupper() or next_line[0].isdigit())


def _text_to_blocks(text: str, page: int | None) -> list[Block]:
    blocks: list[Block] = []
    for para in re.split(r"\n\s*\n", text):
        lines = [ln.rstrip() for ln in para.splitlines() if ln.strip()]
        if not lines:
            continue
        # Markdown heading or a short standalone title line
        if len(lines) == 1 and (lines[0].lstrip().startswith("#") or _looks_like_heading(lines[0].strip())):
            blocks.append(Block(_clean(lines[0].lstrip("# ")), page, "heading"))
            continue
        # "Heading\nBody text..." with no blank line between: peel the heading off.
        if len(lines) > 1 and _is_title_line(lines[0].strip(), lines[1].strip()):
            blocks.append(Block(_clean(lines[0]), page, "heading"))
            lines = lines[1:]
        if all(_LIST_RE.match(ln) for ln in lines):
            for ln in lines:
                blocks.append(Block(_clean(_LIST_RE.sub("", ln)), page, "list"))
            continue
        # "Key: value" lines (syllabus header, resume contact lines) stay one per block
        if len(lines) > 1 and all(len(ln.split()) <= 14 for ln in lines) and sum(":" in ln for ln in lines) >= len(lines) // 2:
            for ln in lines:
                blocks.append(Block(_clean(ln), page, "paragraph"))
            continue
        blocks.append(Block(_clean(" ".join(lines)), page, "paragraph"))
    return blocks


# ---------------------------------------------------------------- images (OCR)
def _load_image(data: bytes, filename: str) -> Document:
    from PIL import Image

    img = Image.open(io.BytesIO(data))
    doc = Document(name=filename, pages=1)
    text = _ocr_image(img)
    if text is None:
        doc.needs_ocr = True
    else:
        doc.blocks.extend(_text_to_blocks(text, 1))
    return doc
